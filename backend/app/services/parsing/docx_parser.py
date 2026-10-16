import docx
from typing import Tuple, List

def extract_text_from_docx(file_path: str) -> Tuple[str, List[str]]:
    """
    Extracts text from a Microsoft Word .docx file including paragraphs and tables.
    Returns: (extracted_text, warnings)
    """
    warnings: List[str] = []
    text_parts: List[str] = []

    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_parts.append(p.text.strip())

        # Also extract table cells
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)

        full_text = "\n\n".join(text_parts).strip()
        if not full_text:
            warnings.append("DOCX file was read successfully but contained no extractable text.")
        return full_text, warnings

    except Exception as e:
        warnings.append(f"python-docx error reading file: {str(e)}")
        return "", warnings
